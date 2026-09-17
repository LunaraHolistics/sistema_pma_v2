# -*- coding: utf-8 -*-
# -*- coding: utf-8 -*-
r"""IMPORTADOR POR PASTA (v4 — Supabase + aba inteligente + flag --gravar).
Uso:
  python importar_pastas.py             -> modo diagnóstico (NÃO grava)
  python importar_pastas.py --gravar    -> grava no Supabase
  python importar_pastas.py --gravar "D:\outro\caminho"
"""
import json, re, os, sys, unicodedata
from datetime import datetime
from openpyxl import load_workbook
import campos, ajustes
import db as DB

RAIZ_PADRAO = r"D:\sistema_pma\dados_importacao"
EXTENSOES = (".xlsx", ".xlsm", ".xlsb")

PASTAS = {
    "cras cecap": ("CRAS Cecap", "cras"), "cras central": ("CRAS Central", "cras"),
    "cras cruzeiro": ("CRAS Cruzeiro do Sul", "cras"), "cras hortencias": ("CRAS Hortênsias", "cras"),
    "cras maria luiza": ("CRAS Maria Luiza", "cras"), "cras parque sao paulo": ("CRAS Parque São Paulo", "cras"),
    "cras sao rafael": ("CRAS São Rafael", "cras"), "cras selmi dei": ("CRAS Selmi Dei", "cras"),
    "cras vale so sol": ("CRAS Vale do Sol", "cras"), "cras vale do sol": ("CRAS Vale do Sol", "cras"),
    "cras valle verde": ("CRAS Valle Verde", "cras"), "cras yolanda opice": ("CRAS Yolanda Ópice", "cras"),
    "creas": ("CREAS", "creas"), "promaip m": ("PROMAIP Masculino", "promaip"),
    "promaip": ("PROMAIP Feminino", "promaip"), "casa de acolhida": ("Casa de Acolhida", "acolhida"),
    "centro dia": ("Centro Dia do Idoso", "centro_dia_idoso"), "centro pop": ("Centro Pop", "centro_pop"),
    "recanto feliz": ("República Idosos Recanto Feliz", "republica_idosos"),
    "residencia inclusiva f": ("Residência Inclusiva", "residencia_inclusiva"),
    "residencia inclusiva": ("Residência Inclusiva", "residencia_inclusiva"),
    "vila dignidade": ("República Idosos Vila Dignidade", "republica_idosos"),
}
MESES_ABREV = ["jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez"]
MESES_COMP  = ["janeiro","fevereiro","marco","abril","maio","junho",
               "julho","agosto","setembro","outubro","novembro","dezembro"]

def norm(t):
    t = unicodedata.normalize("NFKD", str(t).lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"[^a-z0-9]+", " ", t).strip()
    t = re.sub(r"\s+", " ", t)
    return t

def norm_pasta(nome):
    toks = norm(nome).split()
    if toks and toks[0] == "e": toks = toks[1:]
    return " ".join(toks)

def para_int(v):
    if v is None: return None
    if isinstance(v, (int, float)): return int(v)
    s = str(v).strip().upper()
    if s in ("X", "SIM"): return 1
    if s in ("", "-", "NAO", "NÃO"): return 0
    s = re.sub(r"[^\d.-]", "", s)
    try: return int(float(s))
    except: return None

def celulas_texto(ws):
    for r in range(1, min(25, ws.max_row + 1)):
        for c in range(1, min(12, ws.max_column + 1)):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str): yield v

def buscar_ano(textos):
    for t in textos:
        m = re.search(r"20\d{2}", norm(t))
        if m: return int(m.group(0))
    return ajustes.ANO_PADRAO

def detectar_mes_ano(nome, ws):
    n = norm(nome)
    for chave, (m, a) in ajustes.AJUSTES_MANUAIS.items():
        if chave in n: return m, a
    textos = [nome] + list(celulas_texto(ws))
    for i in range(12):
        if MESES_COMP[i] in n: return i + 1, buscar_ano(textos)
    for i in range(12):
        if MESES_ABREV[i] in n.split(): return i + 1, buscar_ano(textos)
    for t in textos:
        tt = norm(t)
        m = re.search(r"20\d{2}", tt)
        if not m: continue
        for i in range(12):
            if MESES_COMP[i] in tt or MESES_ABREV[i] in tt.split():
                return i + 1, int(m.group(0))
    return None, None

def campos_do_tipo(tipo):
    secs = [(s, list(l)) for s, l in campos.CAMPOS[tipo]["secoes"]]
    extra = ajustes.CAMPOS_EXTRAS.get(tipo, [])
    if extra: secs.append(("Complemento", list(extra)))
    return secs

def construir_indice(tipo):
    idx = {}
    for _, lista in campos_do_tipo(tipo):
        for k, r in lista:
            idx.setdefault(norm(r), []).append(k)
    return idx

def match(alvo, idx, pointers):
    if alvo in idx:
        lista = idx[alvo]; p = pointers.get(alvo, 0)
        if p >= len(lista): return None
        pointers[alvo] = p + 1
        return lista[p]
    for nl, lista in idx.items():
        if len(nl) > 8 and (nl in alvo or alvo in nl):
            p = pointers.get(nl, 0)
            if p >= len(lista): continue
            pointers[nl] = p + 1
            return lista[p]
    return None

def chave_do_rotulo(nr):
    if nr in ajustes.ALIASES: return ajustes.ALIASES[nr]
    for pref, k in ajustes.ALIASES_PREFIX:
        if nr.startswith(pref): return k
    return None

def ler_formulario(ws, tipo):
    idx, pointers = construir_indice(tipo), {}
    dados, nao = {}, []
    for r in range(1, ws.max_row + 1):
        partes, valor = [], None
        for c in range(1, min(ws.max_column, 8) + 1):
            v = ws.cell(row=r, column=c).value
            if v is None or isinstance(v, bool) or isinstance(v, datetime): continue
            if isinstance(v, (int, float)): valor = v; break
            s = str(v).strip()
            if not s: continue
            if "ERROR" in s.upper() or "#" in s: break
            if s.upper() in ("X", "SIM", "NÃO", "NAO"): valor = s; break
            partes.append(s)
            if len(partes) == 2: break
        if not partes or valor is None: continue
        rot = " ".join(partes); nr = norm(rot)
        if any(nr.startswith(p) for p in ajustes.IGNORE_PREFIX): continue
        key = chave_do_rotulo(nr) or match(nr, idx, pointers)
        if key: dados[key] = para_int(valor)
        else: nao.append(rot[:60])
    return dados, nao

def processar_arquivo(caminho, arq, tipo):
    """Testa todas as abas e retorna a que casar mais campos."""
    wb = load_workbook(caminho, data_only=True)
    melhor = None
    for nome_aba in wb.sheetnames:
        ws = wb[nome_aba]
        mes, ano = detectar_mes_ano(arq, ws)
        if not mes: continue
        dados, nao = ler_formulario(ws, tipo)
        if melhor is None or len(dados) > len(melhor[0]):
            melhor = (dados, nao, mes, ano, nome_aba)
    return melhor

def main():
    gravar = "--gravar" in sys.argv
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    raiz = args[0] if args else RAIZ_PADRAO
    print(f"📂 Raiz: {raiz}")
    print("🔍 MODO DIAGNÓSTICO — nada será gravado\n" if not gravar else "💾 MODO GRAVAÇÃO (Supabase)\n")
    linhas, prontos, pendencias, vistos = [], [], [], set()

    for nome_pasta in sorted(os.listdir(raiz)):
        caminho_pasta = os.path.join(raiz, nome_pasta)
        if not os.path.isdir(caminho_pasta): continue
        chave = norm_pasta(nome_pasta)
        if chave not in PASTAS:
            pendencias.append(f"PASTA DESCONHECIDA: {nome_pasta}"); continue
        unidade, tipo = PASTAS[chave]
        if tipo not in campos.CAMPOS:
            pendencias.append(f"SEM DNA: {nome_pasta} (tipo '{tipo}')"); continue

        for arq in sorted(os.listdir(caminho_pasta)):
            if arq.startswith("~$"): continue
            if not arq.lower().endswith(EXTENSOES): continue
            caminho = os.path.join(caminho_pasta, arq)
            try:
                melhor = processar_arquivo(caminho, arq, tipo)
            except Exception as e:
                pendencias.append(f"ERRO AO ABRIR: {nome_pasta}/{arq} ({e})"); continue
            if not melhor:
                pendencias.append(f"MÊS/ANO NÃO DETECTADO: {nome_pasta}/{arq}"); continue
            dados, nao, mes, ano, aba = melhor
            dup = (unidade, mes, ano)
            if dup in vistos:
                pendencias.append(f"DUPLICADO: {nome_pasta}/{arq}"); continue
            vistos.add(dup)

            total = sum(len(l) for _, l in campos_do_tipo(tipo))
            status = "✅" if len(dados) >= 5 else "⚠️"
            if len(dados) < 5:
                pendencias.append(f"POUCOS CAMPOS CASADOS: {nome_pasta}/{arq} ({len(dados)})")
            linhas.append(f"{status} {nome_pasta} / {arq}  [aba: {aba}]")
            linhas.append(f"   unidade={unidade} | tipo={tipo} | mês={mes} | ano={ano} | campos casados={len(dados)}/{total}")
            if nao: linhas.append(f"   não casados ({len(nao)}): {', '.join(nao[:6])}")
            prontos.append((unidade, tipo, mes, ano, dados))

    rel = "\n".join(linhas)
    resumo = (f"\n{'='*60}\n✅ PRONTOS PARA GRAVAR: {len(prontos)}\n⚠️ PENDÊNCIAS: {len(pendencias)}\n"
              + "\n".join(f"   - {p}" for p in pendencias))
    print(rel + resumo)
    with open("relatorio_importacao.txt", "w", encoding="utf-8") as f:
        f.write(rel + resumo)
    print("\n💾 Relatório completo em relatorio_importacao.txt")

    if gravar and prontos:
        conn = DB.get_conn()
        for unidade, tipo, mes, ano, dados in prontos:
            DB.upsert_relatorio(conn, unidade, tipo, mes, ano, "importacao_pastas",
                                json.dumps(dados, ensure_ascii=False))
        conn.close()
        print(f"💾 {len(prontos)} registros gravados no Supabase!")
    elif gravar:
        print("⚠️ Nada para gravar.")

if __name__ == '__main__':
    main()